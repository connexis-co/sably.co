#!/usr/bin/env python3
"""Genera docs/data/acortadores-hotmart.xlsx: el mapa de trabajo de los acortadores.

Cruza tres fuentes:
  - src/content/courses/*.mdx        → catálogo real publicado (slug, título, categoría)
  - docs/data/seleccion-cursos.json  → archiveSlug = slug del producto en MasterClasses.LA
  - docs/data/hotmart-map.json       → productos ya identificados con su ID

La nomenclatura del acortador se deriva del archiveSlug, convención verificada contra
un acortador real ya existente el 2026-08-10:
    barberia-artistica-paso-a-paso-curso-crashing  -> pay.hotmart.com/S63192888Y?ref=X76206206C
    barberia-artistica-paso-a-paso-curso-venta-SO  -> ...?checkoutMode=10&ref=X76206206C
"""

from __future__ import annotations

import json
import pathlib
import re
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent
COURSES = ROOT / "src/content/courses"

COLUMNS = [
    ("categoria", 20),
    ("subcategoria", 20),
    ("slug_sably", 42),
    ("titulo_curso", 46),
    ("archive_slug", 46),
    ("id_producto_hotmart", 20),
    ("url_crashing", 58),
    ("url_venta_SO", 58),
    ("ref_afiliado", 16),
    ("aplica_descuento", 17),
    ("codigo_offDiscount", 19),
    ("url_final_publicada", 58),
    ("estado_afiliacion", 19),
    ("verificado_http", 16),
    ("atribucion_ok", 15),
    ("notas", 50),
]


def frontmatter(text: str) -> str:
    m = re.match(r"^---\n(.*?)\n---", text, re.S)
    return m.group(1) if m else ""


def campo(bloque: str, clave: str) -> str:
    m = re.search(rf"^{clave}:\s*(.+)$", bloque, re.M)
    return m.group(1).strip().strip("'\"") if m else ""


def main() -> None:
    seleccion = json.loads((ROOT / "docs/data/seleccion-cursos.json").read_text("utf-8"))
    mapa = json.loads((ROOT / "docs/data/hotmart-map.json").read_text("utf-8"))

    por_sably = {c["sablySlug"]: c for c in seleccion}
    conocidos: dict[str, dict] = {}
    for bloque in ("existentes", "por_crear"):
        conocidos.update(mapa.get(bloque, {}))

    filas: list[dict] = []
    for path in sorted(COURSES.glob("*.mdx")):
        fm = frontmatter(path.read_text("utf-8"))
        slug = path.stem
        url_actual = campo(fm, "hotmartUrl")
        pendiente = "PENDIENTE" in url_actual
        sel = por_sably.get(slug, {})
        archive = sel.get("archiveSlug", "")
        prod = conocidos.get(slug, {})

        notas = []
        if not pendiente:
            notas.append("Ya tenía URL; falta el ref de afiliado")
        if not archive:
            notas.append("Sin archiveSlug: no proviene de MasterClasses.LA, verificar que el producto exista")

        filas.append(
            {
                "categoria": campo(fm, "category"),
                "subcategoria": campo(fm, "subcategory"),
                "slug_sably": slug,
                "titulo_curso": campo(fm, "title"),
                "archive_slug": archive,
                "id_producto_hotmart": prod.get("id", ""),
                "url_crashing": f"https://hotm.art/{archive}-curso-crashing" if archive else "",
                "url_venta_SO": f"https://hotm.art/{archive}-curso-venta-SO" if archive else "",
                "ref_afiliado": "",
                "aplica_descuento": "SI",
                "codigo_offDiscount": "031016",
                "url_final_publicada": "" if pendiente else url_actual,
                "estado_afiliacion": "GLOBAL (Mauricio Duque)" if archive else "POR VERIFICAR",
                "verificado_http": "NO",
                "atribucion_ok": "NO",
                "notas": " · ".join(notas),
            }
        )

    # Plan de trabajo para scripts/hotmart-acortadores.mjs. El nombre del producto
    # se reconstruye del archiveSlug porque es como aparece en "Soy Afiliado(a)".
    plan = [
        {
            "slugSably": f["slug_sably"],
            "titulo": f["titulo_curso"],
            "nombreProducto": f["archive_slug"].replace("-", " ").title(),
            "slugCrashing": f"{f['archive_slug']}-curso-crashing",
            "slugVentaSO": f"{f['archive_slug']}-curso-venta-SO",
        }
        for f in filas
        if f["archive_slug"] and not f["url_final_publicada"]
    ]
    (ROOT / "docs/data/acortadores-plan.json").write_text(
        json.dumps(plan, ensure_ascii=False, indent=2), "utf-8"
    )

    # Resultado real de la corrida del script, si ya existe: manda sobre lo previsto.
    resultado_path = ROOT / "docs/data/acortadores-resultado.json"
    por_slug: dict[str, dict] = {}
    if resultado_path.exists():
        for r in json.loads(resultado_path.read_text("utf-8")):
            por_slug[r["slugSably"]] = r

    for fila in filas:
        r = por_slug.get(fila["slug_sably"])
        if not r:
            continue
        fila["id_producto_hotmart"] = r.get("idProducto") or fila["id_producto_hotmart"]
        fila["url_crashing"] = r.get("crashing") or fila["url_crashing"]
        fila["url_final_publicada"] = r.get("crashing") or fila["url_final_publicada"]
        fila["estado_afiliacion"] = r.get("estado", fila["estado_afiliacion"])
        fila["verificado_http"] = "SI" if r.get("estado") == "OK" else "NO"
        fila["atribucion_ok"] = "SI" if r.get("estado") == "OK" else "NO"
        if r.get("nota"):
            fila["notas"] = (fila["notas"] + " · " if fila["notas"] else "") + r["nota"]

    wb = Workbook()
    wb.remove(wb.active)

    cabecera = Font(bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor="1F2937")
    # Rojo = no se pudo crear el acortador o no quedó verificado. Es la lista de
    # cursos que hay que resolver a mano o sustituyendo el producto.
    rojo = PatternFill("solid", fgColor="FFC7CE")
    rojo_txt = Font(color="9C0006")
    verde = PatternFill("solid", fgColor="C6EFCE")
    verde_txt = Font(color="006100")

    def problematica(f: dict) -> bool:
        return f["verificado_http"] != "SI" or f["atribucion_ok"] != "SI"

    def hoja(nombre: str, datos: list[dict]) -> None:
        ws = wb.create_sheet(nombre[:31])
        ws.append([c for c, _ in COLUMNS])
        for celda in ws[1]:
            celda.font = cabecera
            celda.fill = relleno
            celda.alignment = Alignment(vertical="center")
        for fila in datos:
            ws.append([fila[c] for c, _ in COLUMNS])
            pinta, fuente = (rojo, rojo_txt) if problematica(fila) else (verde, verde_txt)
            for celda in ws[ws.max_row]:
                celda.fill = pinta
                celda.font = fuente
        for i, (_, ancho) in enumerate(COLUMNS, start=1):
            ws.column_dimensions[get_column_letter(i)].width = ancho
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    hoja("RESUMEN", sorted(filas, key=lambda f: (f["categoria"], f["titulo_curso"])))

    por_categoria: dict[str, list[dict]] = defaultdict(list)
    for fila in filas:
        por_categoria[fila["categoria"] or "sin-categoria"].append(fila)
    for categoria in sorted(por_categoria):
        hoja(categoria, sorted(por_categoria[categoria], key=lambda f: f["titulo_curso"]))

    destino = ROOT / "docs/data/acortadores-hotmart.xlsx"
    wb.save(destino)

    con_archive = sum(1 for f in filas if f["archive_slug"])
    print(f"{destino.relative_to(ROOT)}")
    print(f"  {len(filas)} cursos · {len(por_categoria)} categorías")
    print(f"  {con_archive} con archiveSlug (producto de Mauricio Duque)")
    print(f"  {len(filas) - con_archive} sin archiveSlug (verificar manualmente)")


if __name__ == "__main__":
    main()
